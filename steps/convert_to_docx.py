import re
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def set_cell_shading(cell, color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), color)
    tcPr.append(shd)

def parse_flowchart(code_lines):
    steps = []
    for line in code_lines:
        stripped = line.strip()
        if stripped and not stripped.startswith('|') and not stripped.startswith('v'):
            steps.append(stripped)
    return steps

def add_flowchart(doc, steps):
    if not steps:
        return
    table = doc.add_table(rows=0, cols=1)
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    table.autofit = False
    table.allow_autofit = False
    
    tbl = table._tbl
    tblPr = tbl.tblPr if tbl.tblPr is not None else OxmlElement('w:tblPr')
    tblW = OxmlElement('w:tblW')
    tblW.set(qn('w:w'), '6000')
    tblW.set(qn('w:type'), 'pct')
    existing_tblW = tblPr.find(qn('w:tblW'))
    if existing_tblW is not None:
        tblPr.remove(existing_tblW)
    tblPr.append(tblW)
    
    tblBorders = OxmlElement('w:tblBorders')
    for border_name in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        border = OxmlElement(f'w:{border_name}')
        border.set(qn('w:val'), 'single')
        border.set(qn('w:sz'), '4')
        border.set(qn('w:space'), '0')
        border.set(qn('w:color'), '000000')
        tblBorders.append(border)
    existing_borders = tblPr.find(qn('w:tblBorders'))
    if existing_borders is not None:
        tblPr.remove(existing_borders)
    tblPr.append(tblBorders)
    
    for i, step in enumerate(steps):
        row = table.add_row()
        cell = row.cells[0]
        cell.text = ''
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(8)
        run = p.add_run(step)
        run.bold = True
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
        set_cell_shading(cell, 'D9E1F2')
        
        if i < len(steps) - 1:
            arrow_row = table.add_row()
            arrow_cell = arrow_row.cells[0]
            arrow_cell.text = ''
            p = arrow_cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run('▼')
            run.font.size = Pt(14)
            run.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
            set_cell_shading(arrow_cell, 'FFFFFF')

def md_to_docx(md_path, docx_path):
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    doc = Document()
    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(11)

    in_code_block = False
    code_lines = []
    code_lang = ''

    for line in lines:
        stripped = line.rstrip('\n')

        if stripped.startswith('```'):
            if not in_code_block:
                in_code_block = True
                code_lang = stripped[3:].strip()
                code_lines = []
            else:
                in_code_block = False
                if code_lang == 'text' and any('|' in l or 'v' in l for l in code_lines):
                    steps = parse_flowchart(code_lines)
                    add_flowchart(doc, steps)
                else:
                    p = doc.add_paragraph()
                    p.style = 'No Spacing'
                    run = p.add_run(' '.join(code_lines))
                    run.font.name = 'Courier New'
                    run.font.size = Pt(10)
                code_lines = []
            continue

        if in_code_block:
            code_lines.append(stripped)
            continue

        if stripped.startswith('# '):
            doc.add_heading(stripped[2:].strip(), level=1)
        elif stripped.startswith('## '):
            doc.add_heading(stripped[3:].strip(), level=2)
        elif stripped.startswith('### '):
            doc.add_heading(stripped[4:].strip(), level=3)
        elif stripped.startswith('- ') or stripped.startswith('* '):
            doc.add_paragraph(stripped[2:].strip(), style='List Bullet')
        elif re.match(r'^\d+\.\s', stripped):
            doc.add_paragraph(re.sub(r'^\d+\.\s', '', stripped), style='List Number')
        elif stripped == '':
            continue
        else:
            text = stripped
            parts = re.split(r'(\*\*.*?\*\*|\*.*?\*)', text)
            if len(parts) == 1:
                doc.add_paragraph(text)
            else:
                p = doc.add_paragraph()
                for part in parts:
                    if part.startswith('**') and part.endswith('**'):
                        run = p.add_run(part[2:-2])
                        run.bold = True
                    elif part.startswith('*') and part.endswith('*'):
                        run = p.add_run(part[1:-1])
                        run.italic = True
                    else:
                        p.add_run(part)

    doc.save(docx_path)
    print(f"Converted to {docx_path}")

md_to_docx(
    r'D:\vs_projects\freelancing\cale_project\tender-project\PRODUCT_DEMO_OVERVIEW.md',
    r'D:\vs_projects\freelancing\cale_project\tender-project\product_overview.docx'
)
